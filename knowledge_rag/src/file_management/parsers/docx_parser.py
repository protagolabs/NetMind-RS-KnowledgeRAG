"""
DOCX Parser - Handles Microsoft Word documents using python-docx
"""

from pathlib import Path
from typing import List, Dict, Any
import logging

from .base_parser import BaseParser
from ..models import FileMetadata, Episode, ParsingResult, FileFormat


class DocxParser(BaseParser):
    """Parser for Microsoft Word documents"""
    
    SUPPORTED_FORMATS = {FileFormat.DOCX, FileFormat.DOC}
    
    def __init__(self):
        super().__init__()
        self._check_dependencies()
    
    def _check_dependencies(self):
        """Check if required dependencies are available"""
        try:
            import docx
            self.docx = docx
        except ImportError:
            self.logger.error("python-docx not available. Install with: pip install python-docx")
            self.docx = None
    
    def can_parse(self, file_format: FileFormat) -> bool:
        """Check if this parser can handle the given file format"""
        # Note: python-docx only supports .docx, not .doc
        return file_format == FileFormat.DOCX and self.docx is not None
    
    def parse(self, file_path: Path, metadata: FileMetadata) -> ParsingResult:
        """Parse DOCX file and create episodes"""
        if not self.docx:
            return ParsingResult(
                file_id=metadata.file_hash,
                episodes=[],
                success=False,
                error_message="python-docx not available"
            )
        
        if metadata.file_format == FileFormat.DOC:
            return ParsingResult(
                file_id=metadata.file_hash,
                episodes=[],
                success=False,
                error_message="Legacy .doc format not supported. Convert to .docx first."
            )
        
        try:
            episodes = self._parse_docx(file_path, metadata)
            
            return ParsingResult(
                file_id=metadata.file_hash,
                episodes=episodes,
                parsing_metadata={
                    'parser': 'DocxParser',
                    'format': metadata.file_format.value,
                    'total_episodes': len(episodes)
                },
                success=True
            )
            
        except Exception as e:
            self.logger.error(f"Error parsing DOCX file {file_path}: {e}")
            return ParsingResult(
                file_id=metadata.file_hash,
                episodes=[],
                success=False,
                error_message=str(e)
            )
    
    def _parse_docx(self, file_path: Path, metadata: FileMetadata) -> List[Episode]:
        """Parse DOCX document by sections/paragraphs"""
        episodes = []
        
        try:
            doc = self.docx.Document(str(file_path))
            
            # Extract document properties
            doc_properties = self._extract_document_properties(doc)
            
            # Group content by headings/sections
            sections = self._extract_sections(doc)
            
            if sections:
                # Create episodes from sections
                for i, section in enumerate(sections):
                    if section['content'].strip():
                        episode = self.create_episode(
                            content=section['content'],
                            episode_type='section',
                            sequence_number=i + 1,
                            source_file_id=metadata.file_hash,
                            metadata={
                                'heading': section['heading'],
                                'heading_level': section['level'],
                                'has_tables': section['has_tables'],
                                'has_images': section['has_images'],
                                'paragraph_count': section['paragraph_count'],
                                'document_properties': doc_properties
                            }
                        )
                        episodes.append(episode)
            else:
                # Fallback: create episodes from paragraphs
                episodes = self._parse_by_paragraphs(doc, metadata.file_hash, doc_properties)
            
        except Exception as e:
            self.logger.error(f"Error processing DOCX document: {e}")
            raise
        
        return episodes
    
    def _extract_document_properties(self, doc) -> Dict[str, Any]:
        """Extract document metadata and properties"""
        properties = {}
        
        try:
            core_props = doc.core_properties
            
            properties.update({
                'title': core_props.title or '',
                'author': core_props.author or '',
                'subject': core_props.subject or '',
                'created': core_props.created.isoformat() if core_props.created else None,
                'modified': core_props.modified.isoformat() if core_props.modified else None,
                'last_modified_by': core_props.last_modified_by or '',
                'comments': core_props.comments or ''
            })
            
        except Exception as e:
            self.logger.debug(f"Error extracting document properties: {e}")
        
        return properties
    
    def _extract_sections(self, doc) -> List[Dict[str, Any]]:
        """Extract document sections based on headings"""
        sections = []
        current_section = {
            'heading': 'Introduction',
            'level': 0,
            'content': '',
            'has_tables': False,
            'has_images': False,
            'paragraph_count': 0
        }
        
        for paragraph in doc.paragraphs:
            # Check if this is a heading
            if paragraph.style.name.startswith('Heading'):
                # Save previous section if it has content
                if current_section['content'].strip():
                    sections.append(current_section.copy())
                
                # Start new section
                heading_level = int(paragraph.style.name.split()[-1]) if paragraph.style.name.split()[-1].isdigit() else 1
                current_section = {
                    'heading': paragraph.text or f'Heading {len(sections) + 1}',
                    'level': heading_level,
                    'content': paragraph.text + '\n\n',
                    'has_tables': False,
                    'has_images': False,
                    'paragraph_count': 1
                }
            else:
                # Add to current section
                text = paragraph.text
                if text.strip():
                    current_section['content'] += text + '\n\n'
                    current_section['paragraph_count'] += 1
                
                # Check for inline images
                if paragraph.runs:
                    for run in paragraph.runs:
                        if run.element.xpath('.//pic:pic'):
                            current_section['has_images'] = True
        
        # Add the last section
        if current_section['content'].strip():
            sections.append(current_section)
        
        # Check for tables
        for table in doc.tables:
            # Find which section this table belongs to (rough approximation)
            if sections:
                sections[-1]['has_tables'] = True
        
        return sections
    
    def _parse_by_paragraphs(self, doc, file_id: str, doc_properties: Dict) -> List[Episode]:
        """Fallback: parse document by grouping paragraphs"""
        episodes = []
        
        # Group paragraphs into episodes (e.g., every 5-10 paragraphs)
        paragraphs_per_episode = 8
        current_paragraphs = []
        episode_num = 1
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                current_paragraphs.append(text)
                
                if len(current_paragraphs) >= paragraphs_per_episode:
                    content = '\n\n'.join(current_paragraphs)
                    
                    episode = self.create_episode(
                        content=content,
                        episode_type='paragraph_group',
                        sequence_number=episode_num,
                        source_file_id=file_id,
                        metadata={
                            'paragraph_count': len(current_paragraphs),
                            'document_properties': doc_properties,
                            'grouping_method': 'paragraph_based'
                        }
                    )
                    episodes.append(episode)
                    
                    current_paragraphs = []
                    episode_num += 1
        
        # Handle remaining paragraphs
        if current_paragraphs:
            content = '\n\n'.join(current_paragraphs)
            episode = self.create_episode(
                content=content,
                episode_type='paragraph_group',
                sequence_number=episode_num,
                source_file_id=file_id,
                metadata={
                    'paragraph_count': len(current_paragraphs),
                    'document_properties': doc_properties,
                    'grouping_method': 'paragraph_based'
                }
            )
            episodes.append(episode)
        
        # Also extract tables as separate episodes
        table_episodes = self._extract_tables(doc, file_id, len(episodes))
        episodes.extend(table_episodes)
        
        return episodes
    
    def _extract_tables(self, doc, file_id: str, start_sequence: int) -> List[Episode]:
        """Extract tables as separate episodes"""
        table_episodes = []
        
        for table_idx, table in enumerate(doc.tables):
            try:
                # Convert table to text representation
                table_content = self._table_to_text(table, table_idx + 1)
                
                if table_content.strip():
                    episode = self.create_episode(
                        content=table_content,
                        episode_type='table',
                        sequence_number=start_sequence + table_idx + 1,
                        source_file_id=file_id,
                        metadata={
                            'table_index': table_idx + 1,
                            'rows': len(table.rows),
                            'columns': len(table.columns) if table.rows else 0,
                            'content_type': 'table'
                        }
                    )
                    table_episodes.append(episode)
                    
            except Exception as e:
                self.logger.debug(f"Error extracting table {table_idx}: {e}")
        
        return table_episodes
    
    def _table_to_text(self, table, table_num: int) -> str:
        """Convert table to readable text format"""
        content = f"Table {table_num}:\n\n"
        
        try:
            rows_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    row_data.append(cell_text)
                rows_data.append(row_data)
            
            if not rows_data:
                return content + "Empty table\n"
            
            # Format as simple text table
            if len(rows_data) > 0:
                # Use first row as headers if available
                headers = rows_data[0]
                content += "Headers: " + " | ".join(headers) + "\n\n"
                
                # Add data rows
                for i, row_data in enumerate(rows_data[1:], 1):
                    content += f"Row {i}:\n"
                    for j, (header, value) in enumerate(zip(headers, row_data)):
                        content += f"  {header}: {value}\n"
                    content += "\n"
            
        except Exception as e:
            self.logger.debug(f"Error formatting table: {e}")
            content += "Error formatting table content\n"
        
        return content 